from pathlib import Path
import json, re, hashlib, subprocess, os, sysconfig
from collections import Counter
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parents[1]
OUTPUT = BASE / 'MENTOR_REPORT_word.docx'
ENV = dict(os.environ, PYTHONUTF8='1', PYTHONIOENCODING='utf-8')
subprocess.run([str(Path(sysconfig.get_path('scripts')) / ('mddocx.exe' if os.name == 'nt' else 'mddocx')), str(BASE / 'MENTOR_REPORT.md'), '-o', str(OUTPUT), '--theme', 'academic', '--page-size', 'A4', '--font', 'Times New Roman', '--heading-font', 'Times New Roman', '--east-asia-font', '宋体', '--page-numbers', '--strict-math', '--strict-images', '--block-remote-resources', '--diagnostics-json', str(BASE / 'assets/docx_conversion_diagnostics.json')], cwd=BASE, env=ENV, check=True)
doc = Document(OUTPUT)

def fonts(style, size, chinese='宋体', bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = bold
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), chinese)

for name in ['Normal', 'MD Normal']:
    fonts(doc.styles[name], 12)
    pf = doc.styles[name].paragraph_format
    pf.line_spacing = 1.4
    pf.space_after = Pt(6)
    pf.widow_control = True
fonts(doc.styles['Title'], 18, '黑体', True)
fonts(doc.styles['Heading 1'], 15, '黑体', True)
fonts(doc.styles['Heading 2'], 13, '黑体', True)
fonts(doc.styles['MD Equation'], 12)
for name in ['Heading 1', 'Heading 2']:
    pf = doc.styles[name].paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    pf.page_break_before = False

for sec in doc.sections:
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)
    sec.header_distance = sec.footer_distance = Cm(1.2)
    for p in sec.footer.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

for i, p in enumerate(doc.paragraphs):
    oldstyle = p.style.name
    pf = p.paragraph_format
    if i == 0:
        p.style = doc.styles['Title']
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_after = Pt(10)
        pf.keep_with_next = True
    elif oldstyle == 'MD Heading 2':
        p.style = doc.styles['Heading 1']
    elif oldstyle == 'MD Heading 3':
        p.style = doc.styles['Heading 2']
    elif oldstyle == 'MD Normal':
        pf.first_line_indent = Pt(24)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if i == 1 or p.text == '摘要':
        pf.first_line_indent = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.keep_with_next = True
    if p.text.startswith(('关键词：', 'deployment_approved=', 'real_fault_performance_verified=')):
        pf.first_line_indent = Pt(0)
    if re.match(r'^(表|图)\d', p.text):
        pf.first_line_indent = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.line_spacing = 1.15
        pf.space_before, pf.space_after = Pt(5), Pt(7)
        pf.keep_with_next = p.text.startswith('表')
    if p._p.xpath('./w:pPr/w:numPr'):
        pf.first_line_indent = None
    if oldstyle == 'MD Equation':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.line_spacing = 1.0
        pf.space_before = Pt(6)
        pf.space_after = Pt(0)
        pf.keep_with_next = True
    if re.fullmatch(r'（\d+）', p.text):
        pf.first_line_indent = Pt(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if p._p.xpath('.//w:drawing'):
        pf.space_before = Pt(12)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Pt(0)
        pf.keep_with_next = True
        pf.space_after = Pt(0)
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体' if p.style.name in ['Title', 'Heading 1', 'Heading 2'] else '宋体')
        r.font.color.rgb = RGBColor(0, 0, 0)
        if re.match(r'^(表|图)\d', p.text): r.font.size = Pt(10.5)
        elif p.style.name == 'Title': r.font.size = Pt(18)
        elif p.style.name == 'Heading 1': r.font.size = Pt(15)
        elif p.style.name == 'Heading 2': r.font.size = Pt(13)
        else: r.font.size = Pt(12)

for shape in doc.inline_shapes:
    ratio = shape.height / shape.width
    shape.width = Cm(16)
    shape.height = int(shape.width * ratio)

for ti, table in enumerate(doc.tables):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [[3.3, 4.3, 2.2, 6.2], [4.0, 2.4, 2.4, 2.4, 2.4, 2.4], [3.0, 6.5, 6.5]][ti]
    for col, width in zip(table.columns, widths): col.width = Cm(width)
    for ri, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = ri < len(table.rows) - 1
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(10.5)
                    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
                    r.font.bold = ri == 0
                    r.font.color.rgb = RGBColor(0, 0, 0)


# Normalize OOXML constructs that Word handles differently from the converter.
for element in [doc.styles['Title'].element, doc.paragraphs[0]._p]:
    for border in element.xpath('.//w:pBdr'):
        border.getparent().remove(border)

numbering = doc.part.numbering_part.element
abstracts = list(numbering.xpath('./w:abstractNum'))
instances = list(numbering.xpath('./w:num'))
# OOXML requires every abstractNum before every concrete num definition.
for element in abstracts + instances:
    numbering.remove(element)
for element in abstracts:
    if element.find(qn('w:nsid')) is None:
        nsid = OxmlElement('w:nsid')
        nsid.set(qn('w:val'), f'{int(element.get(qn("w:abstractNumId")))+4096:08X}')
        element.insert(0, nsid)
    numbering.append(element)
for element in instances:
    numbering.append(element)
    if element.get(qn('w:numId')) == '11':
        override = OxmlElement('w:lvlOverride')
        override.set(qn('w:ilvl'), '0')
        start = OxmlElement('w:startOverride')
        start.set(qn('w:val'), '1')
        override.append(start)
        element.append(override)

# Render matrix and cases fences as scalable OMML delimiters, not baseline glyphs.
for matrix in list(doc._element.xpath('//m:m')):
    parent = matrix.getparent()
    previous, following = matrix.getprevious(), matrix.getnext()
    def glyph(node):
        if node is None or node.tag != qn('m:r'): return ''
        return ''.join(node.itertext()).strip()
    opening, closing = glyph(previous), glyph(following)
    if opening not in ['[', '{', '(']: continue
    if closing not in [']', '}', ')']: closing = ''
    delimiter = OxmlElement('m:d')
    properties = OxmlElement('m:dPr')
    for tag, value in [('begChr', opening), ('endChr', closing), ('grow', '1')]:
        item = OxmlElement('m:' + tag)
        item.set(qn('m:val'), value)
        properties.append(item)
    delimiter.append(properties)
    expression = OxmlElement('m:e')
    position = parent.index(previous)
    parent.remove(previous)
    if closing: parent.remove(following)
    parent.remove(matrix)
    expression.append(matrix)
    delimiter.append(expression)
    parent.insert(position, delimiter)
for run in doc._element.xpath('//m:r'):
    value = ''.join(run.itertext()).strip()
    if value in ['median', 'std', 'max']:
        props = run.find(qn('m:rPr'))
        if props is None:
            props = OxmlElement('m:rPr')
            run.insert(0, props)
        style = OxmlElement('m:sty')
        style.set(qn('m:val'), 'p')
        props.append(style)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            props = cell._tc.get_or_add_tcPr()
            old = props.find(qn('w:tcBorders'))
            if old is not None: props.remove(old)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                line = OxmlElement('w:' + edge)
                for key, val in [('val', 'single'), ('sz', '4'), ('color', '555555')]:
                    line.set(qn('w:' + key), val)
                borders.append(line)
            props.append(borders)

# Validate that source math was converted to editable OMML, not literal TeX.
md = (BASE / 'MENTOR_REPORT.md').read_text(encoding='utf-8')
inline_count = len(re.findall(r'(?<!\$)\$(?!\$)([^$\n]+)\$(?!\$)', md))
display_count = len(re.findall(r'^\$\$$', md, re.M)) // 2
plain = ''.join(doc._element.xpath('//w:t/text()'))
math_text = ''.join(doc._element.xpath('//m:t/text()'))
assert len(doc.tables) == 3
assert len(doc.inline_shapes) == 3
assert len(doc._element.xpath('//m:oMathPara')) == display_count == 10
assert len(doc._element.xpath('//m:oMath')) == inline_count + display_count
assert '$' not in plain + math_text
assert '**' not in plain
assert not re.search(r'\\(?:noindent|math|begin|operatorname|tau|med|Std)', plain)
assert '下一阶段应优先完成冻结方案' not in plain
counts = Counter(p.style.name for p in doc.paragraphs)
assert counts['Title'] == 1 and counts['Heading 1'] == 7 and counts['Heading 2'] == 16
assert all(not ''.join(sec.header._element.xpath('.//w:t/text()')).strip() for sec in doc.sections)
assert all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for tab in doc.tables for row in tab.rows for cell in row.cells for p in cell.paragraphs)
doc.save(OUTPUT)
audit = dict(output=str(OUTPUT), source_sha256=hashlib.sha256((BASE/'MENTOR_REPORT.tex').read_bytes()).hexdigest(), tables=3, figures=3, inline_equations=inline_count, display_equations=display_count, headings=dict(counts), plain_text_formula_markers=False, header_present=False)
(BASE/'assets/word_structure_audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(audit, ensure_ascii=False))
